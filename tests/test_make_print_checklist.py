"""Tests for the Excel print-tracking checklist."""

import json

import pytest
from openpyxl import load_workbook

from make_print_checklist import build_inventory, main, read_existing_progress


def _roles():
    return {
        "Imp": {
            "image": "https://example.com/imp.png",
            "color": "red",
            # Three copies of the same reminder: one row, quantity three.
            "reminders": ["Dead", "Dead", "Dead"],
            "remindersGlobal": [],
        },
        "Chef": {"image": "u", "color": "blue", "reminders": [], "remindersGlobal": []},
        "Beacon": {
            "image": "u",
            "color": "#FF8800",
            "reminders": ["Watched"],
            "remindersGlobal": ["Global"],
            "_source": "custom",
        },
    }


def test_build_inventory_collapses_duplicate_reminders():
    rows = build_inventory(_roles(), {})

    dead = [r for r in rows if r["role"] == "Imp" and r["type"] == "Reminder"]
    assert len(dead) == 1
    assert dead[0]["qty"] == 3


def test_build_inventory_covers_characters_reminders_and_globals():
    rows = build_inventory(_roles(), {})

    assert sum(1 for r in rows if r["type"] == "Character") == 3
    beacon = {r["reminder"] for r in rows if r["role"] == "Beacon"}
    assert beacon == {"", "Watched", "Global"}  # global reminders included


def test_build_inventory_marks_source_and_script():
    rows = build_inventory(_roles(), {"Homebrew": ["Beacon"]})

    beacon = next(r for r in rows if r["role"] == "Beacon")
    chef = next(r for r in rows if r["role"] == "Chef")
    assert beacon["source"] == "Custom" and beacon["script"] == "Homebrew"
    assert chef["source"] == "Official" and chef["script"] == ""


def test_build_inventory_filenames_match_the_generator():
    rows = build_inventory(_roles(), {})

    character = next(r for r in rows if r["role"] == "Imp" and r["type"] == "Character")
    reminder = next(r for r in rows if r["role"] == "Imp" and r["type"] == "Reminder")
    assert character["file"] == "3mf/red_Imp_coin.3mf"
    assert reminder["file"] == "3mf/red_Imp_Dead_reminder.3mf"
    # A "#RRGGBB" colour must not put a '#' in the filename.
    beacon = next(r for r in rows if r["role"] == "Beacon" and r["type"] == "Character")
    assert "#" not in beacon["file"]


@pytest.fixture
def workbook(tmp_path, monkeypatch):
    import make_print_checklist as mpc

    roles_json = tmp_path / "roles.json"
    roles_json.write_text(json.dumps(_roles()))
    monkeypatch.setattr(mpc, "ROLES_JSON", roles_json)
    monkeypatch.setattr(mpc, "CUSTOM_ROLES_JSON", tmp_path / "absent.json")
    monkeypatch.setattr(mpc, "SCRIPTS_JSON", tmp_path / "absent_scripts.json")
    out = tmp_path / "checklist.xlsx"
    main(["--output", str(out)])
    return out


def test_workbook_has_expected_sheets_and_rows(workbook):
    book = load_workbook(workbook)

    assert book.sheetnames == ["Tokens", "Summary", "How to use"]
    tokens = book["Tokens"]
    assert [c.value for c in tokens[2]][:6] == [
        "Script",
        "Role",
        "Source",
        "Team",
        "Token type",
        "Reminder text",
    ]
    assert tokens.max_row == 2 + 6  # 3 characters + 3 reminder rows


def test_status_is_a_formula_not_a_baked_value(workbook):
    tokens = load_workbook(workbook)["Tokens"]

    status = tokens.cell(row=3, column=9).value
    assert isinstance(status, str) and status.startswith("=IF(")


def test_summary_totals_are_formulas(workbook):
    summary = load_workbook(workbook)["Summary"]

    formulas = [
        cell.value
        for row in summary.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    ]
    assert any(f.startswith("=SUM(") for f in formulas)
    assert any("SUMIFS" in f for f in formulas)


def test_merge_carries_progress_across_a_rebuild(workbook, monkeypatch):
    import make_print_checklist as mpc

    book = load_workbook(workbook)
    tokens = book["Tokens"]
    for row in range(3, tokens.max_row + 1):
        if tokens.cell(row=row, column=2).value == "Chef":
            tokens.cell(row=row, column=8).value = 1
            tokens.cell(row=row, column=11).value = "done on plate 2"
    book.save(workbook)

    progress = read_existing_progress(workbook)
    assert ("Chef", "Character", "") in progress

    mpc.main(["--output", str(workbook), "--merge"])
    rebuilt = load_workbook(workbook)["Tokens"]
    carried = [
        (r[1], r[7], r[10])
        for r in rebuilt.iter_rows(min_row=3, values_only=True)
        if r[7]
    ]
    assert ("Chef", 1, "done on plate 2") in carried


def test_merge_from_seeds_a_new_sheet_without_touching_the_source(
    workbook, tmp_path, capsys, monkeypatch
):
    """A script-scoped sheet can be seeded from the master workbook, and the
    master must be left alone."""
    import make_print_checklist as mpc

    book = load_workbook(workbook)
    tokens = book["Tokens"]
    for row in range(3, tokens.max_row + 1):
        if tokens.cell(row=row, column=2).value in ("Chef", "Imp"):
            tokens.cell(row=row, column=8).value = 1
    book.save(workbook)
    before = tokens.max_row

    scripts = tmp_path / "scripts.json"
    scripts.write_text(json.dumps({"Small": ["Imp"]}))
    monkey = tmp_path / "scoped.xlsx"
    monkeypatch.setattr(mpc, "SCRIPTS_JSON", scripts)
    mpc.main(
        [
            "--script",
            "Small",
            "--merge-from",
            str(workbook),
            "--output",
            str(monkey),
        ]
    )

    # Source untouched.
    assert load_workbook(workbook)["Tokens"].max_row == before
    # Scoped sheet holds only the script's roles, with progress carried over.
    scoped = load_workbook(monkey)["Tokens"]
    roles = {r[1] for r in scoped.iter_rows(min_row=3, values_only=True) if r[1]}
    assert roles == {"Imp"}
    printed = [r[7] for r in scoped.iter_rows(min_row=3, values_only=True) if r[1]]
    assert 1 in printed
    # Writing to a separate file is the intended workflow, so the roles that
    # were simply filtered out must not be reported as a problem.
    output = capsys.readouterr().out
    assert "Chef" not in output
    assert "narrowed" not in output


def test_narrowing_the_master_sheet_warns_about_lost_progress(
    workbook, tmp_path, capsys, monkeypatch
):
    import make_print_checklist as mpc

    book = load_workbook(workbook)
    tokens = book["Tokens"]
    for row in range(3, tokens.max_row + 1):
        if tokens.cell(row=row, column=2).value == "Chef":
            tokens.cell(row=row, column=8).value = 1
    book.save(workbook)

    scripts = tmp_path / "scripts.json"
    scripts.write_text(json.dumps({"Small": ["Imp"]}))
    monkeypatch.setattr(mpc, "SCRIPTS_JSON", scripts)
    mpc.main(["--script", "Small", "--merge", "--output", str(workbook)])

    output = capsys.readouterr().out
    assert "Chef" in output
    assert "You have just narrowed" in output
    assert "no longer in this file" in output
